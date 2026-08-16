"""Read-only loader for the frozen human-aligned Gold decision input."""

from __future__ import annotations

import csv
from collections.abc import Mapping
from dataclasses import dataclass
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.schemas import (
    AnnotationRunManifest,
    CommentRecord,
    DatasetSplitManifest,
    DatasetSplitRole,
    EvidenceAtom,
    EvidenceGrade,
    EvidenceRoute,
    ExperienceScope,
    FormalAnnotationIdentity,
    SampleType,
    ScreeningStatus,
    SourceReference,
    SourceType,
)
from src.services.prepared_corpus import canonical_json_bytes, sha256_bytes


DATA_ROOT = Path(__file__).resolve().parents[2] / "data" / "competition" / "screened_v2"
GOLD_LABELS_RELATIVE = Path(
    "annotations/gold_v2/attempt-0002/comparison/frozen_gold_labels.json"
)
GOLD_WORKBOOK_RELATIVE = Path("datasets/GOLD 人工标注对齐.xlsx")
RECORDS_RELATIVE = Path("screened_v2.csv")
SPLIT_RELATIVE = Path("split_manifest.json")
ANNOTATION_MANIFEST_RELATIVE = Path(
    "annotations/gold_v2/attempt-0002/annotation_run_manifest.json"
)
GOLD_COMPARISON_MANIFEST_RELATIVE = Path(
    "annotations/gold_v2/attempt-0002/comparison/gold_comparison_manifest.json"
)
EXPECTED_GOLD_COUNT = 50
EXPECTED_DATASET_VERSION = "screened_v2"
EXPECTED_GOLD_WORKBOOK = "GOLD 人工标注对齐.xlsx"
EXPECTED_RESOLUTION_STATUSES = {"AGREED", "TEAM_CONSENSUS"}


@dataclass(frozen=True, slots=True)
class GoldDecisionInput:
    """Immutable human Gold records and labels consumed by decision assembly."""

    records: tuple[CommentRecord, ...]
    atoms: tuple[EvidenceAtom, ...]
    result_sha256: str
    calibration_sha256: str
    stable_ids_sha256: str
    annotation_identity: FormalAnnotationIdentity
    labels_json_sha256: str | None = None

    def validate(self) -> None:
        if len(self.records) != EXPECTED_GOLD_COUNT or len(self.atoms) != EXPECTED_GOLD_COUNT:
            raise ValueError("Human Gold must contain exactly 50 records and atoms")
        if len(self.records) != len(self.atoms):
            raise ValueError("Human Gold records and atoms must have equal lengths")
        stable_ids = [record.raw_id or record.comment_id for record in self.records]
        if len(set(stable_ids)) != EXPECTED_GOLD_COUNT:
            raise ValueError("Human Gold stable IDs must be unique")
        expected_sha = _stable_ids_sha256(stable_ids)
        if self.stable_ids_sha256 != expected_sha:
            raise ValueError("Human Gold stable IDs SHA-256 does not match records")
        _require_sha(self.result_sha256, "Human Gold result SHA")
        _require_sha(self.calibration_sha256, "Human Gold calibration SHA")
        if self.labels_json_sha256 is not None:
            _require_sha(self.labels_json_sha256, "frozen Gold labels JSON SHA")
        for record, atom in zip(self.records, self.atoms, strict=True):
            if atom.comment_id != record.comment_id:
                raise ValueError("Human Gold atom comment_id does not match record")
            if atom.label_source != "HUMAN_GOLD":
                raise ValueError("Human Gold atom label_source must be HUMAN_GOLD")
            if atom.ai_confidence is not None or atom.explanation is not None:
                raise ValueError("Human Gold atom must not contain AI confidence or explanation")


def _require_sha(value: object, label: str) -> str:
    if not isinstance(value, str) or len(value) != 64:
        raise ValueError(f"{label} must be a lowercase 64-character SHA-256")
    if any(character not in "0123456789abcdef" for character in value):
        raise ValueError(f"{label} must be a lowercase 64-character SHA-256")
    return value


def _regular_file(path: Path, label: str) -> None:
    if path.is_symlink() or not path.is_file():
        raise ValueError(f"{label} must be an existing regular file: {path}")


def _read_json_object(path: Path, label: str) -> dict[str, Any]:
    _regular_file(path, label)
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError(f"{label} is not valid UTF-8 JSON") from exc
    if not isinstance(value, dict):
        raise ValueError(f"{label} must be a JSON object")
    return value


def _file_sha256(path: Path, label: str) -> str:
    _regular_file(path, label)
    return sha256_bytes(path.read_bytes())


def _stable_id(record: CommentRecord) -> str:
    return record.raw_id or record.comment_id


def _stable_ids_sha256(stable_ids: list[str] | tuple[str, ...]) -> str:
    return sha256_bytes(
        canonical_json_bytes(sorted(stable_ids))
    )


def _text(value: object, label: str) -> str:
    if not isinstance(value, str) or not value:
        raise ValueError(f"{label} must be a non-empty string")
    return value


def _optional_text(value: object) -> str | None:
    if value is None or value == "":
        return None
    if not isinstance(value, str):
        raise ValueError("CSV optional text field must be a string")
    return value


def _csv_bool(value: object, label: str) -> bool | None:
    if value in (None, ""):
        return None
    if value == "True":
        return True
    if value == "False":
        return False
    raise ValueError(f"{label} must be True or False")


def _load_records(path: Path) -> dict[str, CommentRecord]:
    _regular_file(path, "screened_v2.csv")
    records: dict[str, CommentRecord] = {}
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {
            "comment_id",
            "raw_id",
            "sample_type",
            "raw_sample_type",
            "raw_content",
            "context_content",
            "source_platform",
            "original_url",
            "platform_url_available",
            "collected_at",
            "screening_status",
            "screening_reason",
            "source_id",
            "source_type",
            "source_ref",
        }
        if reader.fieldnames is None or not required.issubset(reader.fieldnames):
            raise ValueError("screened_v2.csv is missing required source fields")
        for row in reader:
            source = SourceReference(
                source_id=_text(row.get("source_id"), "source_id"),
                source_type=SourceType(_text(row.get("source_type"), "source_type")),
                source_ref=_text(row.get("source_ref"), "source_ref"),
            )
            record = CommentRecord.model_validate(
                {
                    "comment_id": _text(row.get("comment_id"), "comment_id"),
                    "raw_id": _text(row.get("raw_id"), "raw_id"),
                    "sample_type": SampleType(_text(row.get("sample_type"), "sample_type")),
                    "raw_sample_type": _text(row.get("raw_sample_type"), "raw_sample_type"),
                    "raw_content": _text(row.get("raw_content"), "raw_content"),
                    "context_content": _optional_text(row.get("context_content")),
                    "source_platform": _text(row.get("source_platform"), "source_platform"),
                    "original_url": _optional_text(row.get("original_url")),
                    "platform_url_available": _csv_bool(
                        row.get("platform_url_available"), "platform_url_available"
                    ),
                    "collected_at": _optional_text(row.get("collected_at")),
                    "screening_status": ScreeningStatus(
                        _text(row.get("screening_status"), "screening_status")
                    ),
                    "screening_reason": _optional_text(row.get("screening_reason")),
                    "source": source,
                },
                strict=True,
            )
            stable_id = _stable_id(record)
            if stable_id in records:
                raise ValueError(f"screened_v2.csv contains duplicate stable ID: {stable_id}")
            records[stable_id] = record
    return records


def _load_split(path: Path) -> DatasetSplitManifest:
    _regular_file(path, "split_manifest.json")
    try:
        return DatasetSplitManifest.model_validate_json(path.read_bytes(), strict=True)
    except ValueError as exc:
        raise ValueError("split_manifest.json is invalid") from exc


def _load_annotation_identity(
    path: Path, labels: Mapping[str, Any]
) -> tuple[AnnotationRunManifest, FormalAnnotationIdentity]:
    manifest_payload = _read_json_object(path, "Gold annotation_run_manifest.json")
    try:
        manifest = AnnotationRunManifest.model_validate_json(
            path.read_bytes(), strict=True
        )
    except ValueError as exc:
        raise ValueError("Gold annotation_run_manifest.json is invalid") from exc
    fields = (
        "annotation_version",
        "model_id",
        "reasoning_effort",
        "dataset_version",
        "dataset_sha256",
        "prompt_version",
        "prompt_sha256",
        "result_sha256",
    )
    for field in fields:
        if labels.get(field) != getattr(manifest, field):
            raise ValueError(f"frozen Gold {field} does not match annotation manifest")
    return manifest, FormalAnnotationIdentity(
        annotation_version=manifest.annotation_version,
        model_id=manifest.model_id,
        reasoning_effort=manifest.reasoning_effort,
        dataset_version=manifest.dataset_version,
        dataset_sha256=manifest.dataset_sha256,
        prompt_version=manifest.prompt_version,
        prompt_sha256=manifest.prompt_sha256,
        result_sha256=manifest.result_sha256,
    )


def _validate_comparison_manifest(labels_path: Path, labels_sha256: str) -> None:
    comparison_path = labels_path.with_name("gold_comparison_manifest.json")
    if not comparison_path.exists():
        return
    comparison = _read_json_object(comparison_path, "Gold comparison manifest")
    artifacts = comparison.get("artifact_sha256")
    if not isinstance(artifacts, Mapping):
        raise ValueError("Gold comparison manifest artifact_sha256 is missing")
    expected = artifacts.get(labels_path.name)
    if expected is not None and expected != labels_sha256:
        raise ValueError("frozen Gold labels JSON SHA-256 does not match comparison manifest")


def _load_workbook_rows(path: Path, *, sheet_name: str) -> dict[str, dict[str, object]]:
    _regular_file(path, "human Gold workbook")
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"human Gold workbook sheet is missing: {sheet_name}")
    rows = list(workbook[sheet_name].iter_rows(values_only=True))
    if not rows:
        raise ValueError("human Gold workbook is empty")
    headers = [str(value) if value is not None else "" for value in rows[0]]
    required = {
        "evidence_id",
        "raw_id",
        "raw_content",
        "source_platform",
        "route",
        "evidence_grade",
        "experience_scope",
    }
    if not required.issubset(headers):
        raise ValueError("human Gold workbook is missing required columns")
    indexes = {header: index for index, header in enumerate(headers)}
    result: dict[str, dict[str, object]] = {}
    for values in rows[1:]:
        if not any(value not in (None, "") for value in values):
            continue
        row = {header: values[index] if index < len(values) else None for header, index in indexes.items()}
        raw_id = _text(row.get("raw_id"), "human Gold workbook raw_id")
        if raw_id in result:
            raise ValueError(f"human Gold workbook contains duplicate raw_id: {raw_id}")
        result[raw_id] = row
    if len(result) != EXPECTED_GOLD_COUNT:
        raise ValueError("human Gold workbook must contain exactly 50 rows")
    return result


def _validate_label_payload(
    labels_payload: Mapping[str, Any],
    *,
    split_manifest: DatasetSplitManifest,
    workbook_rows: Mapping[str, Mapping[str, object]],
    records: Mapping[str, CommentRecord],
) -> tuple[list[dict[str, Any]], AnnotationRunManifest, FormalAnnotationIdentity]:
    if labels_payload.get("schema_version") != "frozen_gold_labels_v2":
        raise ValueError("frozen Gold labels schema version is invalid")
    if labels_payload.get("dataset_version") != EXPECTED_DATASET_VERSION:
        raise ValueError("frozen Gold dataset_version is invalid")
    if labels_payload.get("label_count") != EXPECTED_GOLD_COUNT:
        raise ValueError("frozen Gold label_count must be exactly 50")
    labels = labels_payload.get("labels")
    if not isinstance(labels, list) or len(labels) != EXPECTED_GOLD_COUNT:
        raise ValueError("frozen Gold labels must contain exactly 50 labels")
    gold_source = labels_payload.get("gold_source")
    if not isinstance(gold_source, Mapping):
        raise ValueError("frozen Gold gold_source is missing")
    if gold_source.get("artifact") != EXPECTED_GOLD_WORKBOOK:
        raise ValueError("frozen Gold workbook identity is invalid")
    if gold_source.get("sheet") != "人工标注":
        raise ValueError("frozen Gold workbook sheet identity is invalid")
    for column in ("route", "evidence_grade", "experience_scope"):
        if column not in (gold_source.get("label_columns") or []):
            raise ValueError("frozen Gold label columns are incomplete")

    if labels_payload.get("dataset_sha256") != split_manifest.dataset_sha256:
        raise ValueError("frozen Gold dataset SHA does not match split manifest")
    if labels_payload.get("dataset_version") != split_manifest.dataset_version:
        raise ValueError("frozen Gold dataset version does not match split manifest")
    manifest, identity = _load_annotation_identity(
        Path(labels_payload["_annotation_manifest_path"]), labels_payload
    )
    if identity.dataset_sha256 != split_manifest.dataset_sha256:
        raise ValueError("Gold annotation dataset SHA does not match split manifest")

    gold_assignments = [
        assignment
        for assignment in split_manifest.assignments
        if assignment.split_role is DatasetSplitRole.GOLD
    ]
    gold_ids = [assignment.raw_id for assignment in gold_assignments]
    if len(gold_ids) != EXPECTED_GOLD_COUNT or len(set(gold_ids)) != EXPECTED_GOLD_COUNT:
        raise ValueError("split_manifest GOLD role must contain exactly 50 unique IDs")
    labels_by_id: dict[str, dict[str, Any]] = {}
    ordered_labels: list[dict[str, Any]] = []
    for label in labels:
        if not isinstance(label, Mapping):
            raise ValueError("frozen Gold labels must be JSON objects")
        evidence_id = _text(label.get("evidence_id"), "frozen Gold evidence_id")
        if evidence_id in labels_by_id:
            raise ValueError("frozen Gold evidence IDs must be unique")
        if label.get("is_frozen") is not True:
            raise ValueError("every frozen Gold label must have is_frozen=true")
        statuses = (
            label.get("route_resolution_status"),
            label.get("evidence_grade_resolution_status"),
            label.get("experience_scope_resolution_status"),
        )
        if any(status == "UNRESOLVED" for status in statuses):
            raise ValueError("frozen Gold labels must not contain UNRESOLVED fields")
        if any(status not in EXPECTED_RESOLUTION_STATUSES for status in statuses):
            raise ValueError("frozen Gold resolution status is invalid")
        final_values = {
            "final_route": EvidenceRoute(_text(label.get("final_route"), "final_route")),
            "final_evidence_grade": EvidenceGrade(
                _text(label.get("final_evidence_grade"), "final_evidence_grade")
            ),
            "final_experience_scope": ExperienceScope(
                _text(label.get("final_experience_scope"), "final_experience_scope")
            ),
        }
        if evidence_id not in set(gold_ids):
            raise ValueError("frozen Gold IDs must exactly match split_manifest GOLD IDs")
        if evidence_id not in workbook_rows or evidence_id not in records:
            raise ValueError(f"frozen Gold ID is missing from source records: {evidence_id}")
        workbook_row = workbook_rows[evidence_id]
        record = records[evidence_id]
        if workbook_row.get("raw_content") != record.raw_content:
            raise ValueError(f"human Gold raw_content does not match screened_v2.csv: {evidence_id}")
        if workbook_row.get("source_platform") != record.source_platform:
            raise ValueError(f"human Gold source_platform does not match screened_v2.csv: {evidence_id}")
        for workbook_field, expected in (
            ("route", final_values["final_route"].value),
            ("evidence_grade", final_values["final_evidence_grade"].value),
            ("experience_scope", final_values["final_experience_scope"].value),
        ):
            if workbook_row.get(workbook_field) != expected:
                raise ValueError(f"human Gold workbook label does not match frozen JSON: {evidence_id}")
        label_value = dict(label)
        label_value.update(final_values)
        labels_by_id[evidence_id] = label_value
        ordered_labels.append(label_value)
    if set(labels_by_id) != set(gold_ids):
        raise ValueError("frozen Gold IDs do not exactly match split_manifest GOLD IDs")
    ai_ids = labels_payload.get("ai_evidence_ids")
    if ai_ids != [label["evidence_id"] for label in ordered_labels]:
        raise ValueError("frozen Gold ai_evidence_ids do not match frozen label order")
    return ordered_labels, manifest, identity


def _build_input(
    *,
    labels_path: Path,
    workbook_path: Path,
    records_path: Path,
    split_path: Path,
    annotation_manifest_path: Path,
) -> GoldDecisionInput:
    labels_payload = _read_json_object(labels_path, "frozen Gold labels")
    labels_sha256 = _file_sha256(labels_path, "frozen Gold labels")
    _validate_comparison_manifest(labels_path, labels_sha256)
    workbook_sha256 = _file_sha256(workbook_path, "human Gold workbook")
    gold_source = labels_payload.get("gold_source")
    if not isinstance(gold_source, Mapping):
        raise ValueError("frozen Gold gold_source is missing")
    if gold_source.get("artifact_sha256") != workbook_sha256:
        raise ValueError("human Gold artifact SHA does not match workbook")
    split_manifest = _load_split(split_path)
    records = _load_records(records_path)
    annotation_manifest = _read_json_object(
        annotation_manifest_path, "Gold annotation_run_manifest.json"
    )
    labels_for_validation = dict(labels_payload)
    labels_for_validation["_annotation_manifest_path"] = str(annotation_manifest_path)
    ordered_labels, manifest, identity = _validate_label_payload(
        labels_for_validation,
        split_manifest=split_manifest,
        workbook_rows=_load_workbook_rows(
            workbook_path, sheet_name=str(gold_source.get("sheet"))
        ),
        records=records,
    )
    if manifest.model_id != "gpt-5.6-sol" or manifest.reasoning_effort != "high":
        raise ValueError("Human Gold annotation identity must be gpt-5.6-sol/high")
    result_sha256 = _require_sha(labels_payload.get("result_sha256"), "Human Gold result SHA")
    atoms: list[EvidenceAtom] = []
    selected_records: list[CommentRecord] = []
    for label in ordered_labels:
        stable_id = str(label["evidence_id"])
        record = records[stable_id]
        if record.source is None:
            raise ValueError(f"Human Gold source is missing: {stable_id}")
        selected_records.append(record)
        atoms.append(
            EvidenceAtom(
                evidence_id=stable_id,
                comment_id=record.comment_id,
                route=label["final_route"],
                experience_scope=label["final_experience_scope"],
                evidence_grade=label["final_evidence_grade"],
                ai_confidence=None,
                explanation=None,
                label_source="HUMAN_GOLD",
                source=record.source,
                source_platform=record.source_platform,
                duplicate_group=record.duplicate_group,
                actual_use=record.actual_use,
            )
        )
    result = GoldDecisionInput(
        records=tuple(selected_records),
        atoms=tuple(atoms),
        result_sha256=result_sha256,
        calibration_sha256=workbook_sha256,
        stable_ids_sha256=_stable_ids_sha256(
            [_stable_id(record) for record in selected_records]
        ),
        annotation_identity=identity,
        labels_json_sha256=labels_sha256,
    )
    result.validate()
    return result


@dataclass(frozen=True, slots=True)
class HumanGoldDecisionInput(GoldDecisionInput):
    """Named loader type for the immutable, human-aligned Gold input."""

    @classmethod
    def load(
        cls,
        *,
        data_root: Path = DATA_ROOT,
        labels_path: Path | None = None,
        workbook_path: Path | None = None,
        records_path: Path | None = None,
        split_manifest_path: Path | None = None,
        annotation_manifest_path: Path | None = None,
    ) -> "HumanGoldDecisionInput":
        root = Path(data_root)
        result = _build_input(
            labels_path=Path(labels_path or root / GOLD_LABELS_RELATIVE),
            workbook_path=Path(workbook_path or root / GOLD_WORKBOOK_RELATIVE),
            records_path=Path(records_path or root / RECORDS_RELATIVE),
            split_path=Path(split_manifest_path or root / SPLIT_RELATIVE),
            annotation_manifest_path=Path(
                annotation_manifest_path or root / ANNOTATION_MANIFEST_RELATIVE
            ),
        )
        return cls(
            records=result.records,
            atoms=result.atoms,
            result_sha256=result.result_sha256,
            calibration_sha256=result.calibration_sha256,
            stable_ids_sha256=result.stable_ids_sha256,
            annotation_identity=result.annotation_identity,
            labels_json_sha256=result.labels_json_sha256,
        )


def load_human_gold_decision_input(
    *, data_root: Path = DATA_ROOT, **paths: Path
) -> HumanGoldDecisionInput:
    return HumanGoldDecisionInput.load(data_root=data_root, **paths)
