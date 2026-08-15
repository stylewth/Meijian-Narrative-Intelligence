from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.schemas import PublicEvidenceCorpusV2, PublicNarrativeEvidence
from src.services.prepared_corpus import canonical_json_bytes, sha256_bytes


PUBLIC_NARRATIVE_EVIDENCE_HEADERS = (
    "证据编号",
    "对象类型",
    "品牌名称",
    "证据性质",
    "证据主题",
    "官方原文",
    "事实或主张摘要",
    "对应产品",
    "目标人群",
    "主要场景",
    "核心价值",
    "内容母题",
    "Campaign/口号",
    "产品支撑",
    "来源类型",
    "来源链接",
    "发布日期",
    "采集日期",
    "当前有效性",
    "证据可信度",
    "备注",
)
_EXPECTED_DATA_ROW_COUNT = 54


def _normalise_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value).strip()


def _normalise_row(row: tuple[Any, ...], row_number: int) -> tuple[str, ...]:
    if len(row) != len(PUBLIC_NARRATIVE_EVIDENCE_HEADERS):
        raise ValueError(
            f"public narrative evidence row {row_number} must have exactly "
            f"{len(PUBLIC_NARRATIVE_EVIDENCE_HEADERS)} columns"
        )
    return tuple(_normalise_cell(value) for value in row)


def load_public_narrative_evidence(path: Path) -> PublicEvidenceCorpusV2:
    """Read and freeze every row of the public brand/competitor evidence workbook."""

    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise ValueError(f"public narrative evidence workbook does not exist: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        worksheet = workbook.active
        worksheet.reset_dimensions()
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise ValueError("public narrative evidence workbook is empty")

    headers = _normalise_row(rows[0], 1)
    if headers != PUBLIC_NARRATIVE_EVIDENCE_HEADERS:
        raise ValueError(
            "public narrative evidence header does not exactly match the frozen 21-column contract"
        )

    data_rows = rows[1:]
    if len(data_rows) != _EXPECTED_DATA_ROW_COUNT:
        raise ValueError(
            f"public narrative evidence row count must be exactly "
            f"{_EXPECTED_DATA_ROW_COUNT}, got {len(data_rows)}"
        )

    raw_items: list[dict[str, str]] = []
    items: list[PublicNarrativeEvidence] = []
    evidence_ids: set[str] = set()
    for row_number, row in enumerate(data_rows, start=2):
        values = _normalise_row(row, row_number)
        raw_fields = dict(zip(PUBLIC_NARRATIVE_EVIDENCE_HEADERS, values, strict=True))
        evidence_id = raw_fields["证据编号"]
        if not evidence_id:
            raise ValueError(f"public narrative evidence row {row_number} has blank evidence_id")
        if evidence_id in evidence_ids:
            raise ValueError(f"public narrative evidence has duplicate evidence_id: {evidence_id}")
        evidence_ids.add(evidence_id)
        raw_items.append(raw_fields)
        items.append(
            PublicNarrativeEvidence(
                evidence_id=evidence_id,
                brand=raw_fields["品牌名称"],
                source_title=raw_fields["证据主题"],
                source_url=raw_fields["来源链接"],
                source_type=raw_fields["来源类型"],
                claim_text=raw_fields["事实或主张摘要"],
                status=raw_fields["当前有效性"],
                confidence=raw_fields["证据可信度"],
                product_support=raw_fields["产品支撑"],
                raw_fields=raw_fields,
            )
        )

    content_sha256 = sha256_bytes(
        canonical_json_bytes(
            {"headers": list(PUBLIC_NARRATIVE_EVIDENCE_HEADERS), "items": raw_items}
        )
    )
    return PublicEvidenceCorpusV2(
        items=items,
        file_sha256=sha256_bytes(workbook_path.read_bytes()),
        content_sha256=content_sha256,
    )
