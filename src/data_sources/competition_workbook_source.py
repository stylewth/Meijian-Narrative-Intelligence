from __future__ import annotations

import hashlib
import math
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from collections.abc import Iterable, Mapping
from typing import Any

from openpyxl import load_workbook

from src.data_validator import ValidatedDataset
from src.field_mapper import FieldMapping
from src.schemas import (
    CommentRecord,
    DataHealthResult,
    DataHealthStatus,
    SampleType,
    ScreeningStatus,
    SourceReference,
    SourceType,
)


WORKSHEET_NAME = "表1：原始语料库"
EXPECTED_HEADERS = (
    "语料编号",
    "样本类型",
    "项目编号",
    "评论对象",
    "原始内容",
    "上下文内容",
    "来源平台",
    "来源链接",
    "发布时间",
    "采集日期",
    "搜索关键词",
    "点赞数",
    "作者匿名编号",
    "真实性类型",
    "采集人员",
    "备注",
    "对象类型",
    "关联清洗记录",
    "关联分析记录",
)
EXPECTED_PLATFORMS = frozenset({"抖音", "小红书", "京东", "淘宝", "B站"})
DATASET_VERSION = "raw_competition_v1"

RAW_TYPE_MAPPING = {
    "梅见产品体验反馈": SampleType.MEIJIAN_FEEDBACK,
    "梅见品牌认知与传播反馈": SampleType.MEIJIAN_FEEDBACK,
    "竞品产品体验反馈": SampleType.COMPETITOR_FEEDBACK,
    "竞品选择与替代理由": SampleType.COMPETITOR_FEEDBACK,
    "饮酒情绪与身体负担": SampleType.DRINKING_EMOTION,
    "低度酒品类态度": SampleType.DRINKING_EMOTION,
    "社交关系与饮酒边界": SampleType.DRINKING_EMOTION,
    "消费场景与使用行为": SampleType.SCENE_NEED,
    "饮法调配与内容共创": SampleType.SCENE_NEED,
    "购买决策与价格反馈": SampleType.SCENE_NEED,
}


@dataclass(frozen=True, slots=True)
class CompetitionSourceManifest:
    dataset_version: str
    source_sha256: str
    imported_at: str
    source_filename: str
    source_sheet: str
    record_count: int

    def to_dict(self) -> dict[str, str | int]:
        return {
            "dataset_version": self.dataset_version,
            "imported_at": self.imported_at,
            "record_count": self.record_count,
            "source_filename": self.source_filename,
            "source_sha256": self.source_sha256,
            "source_sheet": self.source_sheet,
        }


@dataclass(frozen=True, slots=True)
class CompetitionWorkbookImport:
    dataset: ValidatedDataset
    source_manifest: CompetitionSourceManifest


def scalar(value: object) -> str | None:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    text = str(value).strip()
    return text or None


def normalized_source_ref(*, workbook_name: str, row_number: int) -> str:
    return f"{workbook_name}#{WORKSHEET_NAME}!{row_number}"


def normalize_original_url(value: object) -> tuple[str | None, bool]:
    text = scalar(value)
    if text in {None, "数据未收集"}:
        return None, False
    if not text.startswith(("https://", "http://")):
        raise ValueError(f"来源链接不是 HTTP(S) URL: {text}")
    return text, True


def _formatted_date(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return scalar(value)


def _source_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _field_mapping() -> FieldMapping:
    return FieldMapping(
        record_fields={
            "raw_id": "语料编号",
            "raw_sample_type": "样本类型",
            "raw_content": "原始内容",
            "context_content": "上下文内容",
            "source_platform": "来源平台",
            "original_url": "来源链接",
            "collected_at": "采集日期",
            "source_note": "备注",
        },
        annotation_fields={},
    )


def _data_health(comments: list[CommentRecord]) -> DataHealthResult:
    effective_comments = [
        comment
        for comment in comments
        if comment.screening_status is ScreeningStatus.KEEP
        and comment.source is not None
        and comment.source.source_type is SourceType.USER_COMMENT
    ]
    platform_counts: dict[str, int] = {}
    for comment in effective_comments:
        assert comment.source_platform is not None
        platform_counts[comment.source_platform] = platform_counts.get(comment.source_platform, 0) + 1
    record_count = len(effective_comments)
    largest_platform = max(platform_counts, key=platform_counts.get) if platform_counts else None
    largest_share = (
        platform_counts[largest_platform] / record_count
        if largest_platform is not None and record_count
        else 0.0
    )
    traceable_rate = (
        sum(comment.source is not None for comment in effective_comments) / record_count
        if record_count
        else 0.0
    )
    platform_url_coverage = (
        sum(comment.platform_url_available is True for comment in effective_comments) / record_count
        if record_count
        else 0.0
    )
    ready = (
        record_count > 0
        and len(platform_counts) >= 4
        and largest_share <= 0.5
        and traceable_rate == 1.0
    )
    return DataHealthResult(
        status=DataHealthStatus.READY if ready else DataHealthStatus.NEEDS_REBALANCE,
        effective_user_comment_count=record_count,
        platform_count=len(platform_counts),
        largest_platform=largest_platform,
        largest_platform_share=largest_share,
        traceable_rate=traceable_rate,
        platform_url_coverage=platform_url_coverage,
        actual_use_ratio=None,
        simulated_record_count=0,
        can_finalize_snapshot=ready,
    )


def _decision_value(decision: object, field: str) -> object:
    if isinstance(decision, Mapping):
        return decision.get(field)
    return getattr(decision, field, None)


def apply_screening_decisions(
    imported: CompetitionWorkbookImport,
    decisions: Iterable[object],
) -> CompetitionWorkbookImport:
    """Apply one explicit screening decision to every raw record."""

    by_id: dict[str, tuple[ScreeningStatus, str]] = {}
    for decision in decisions:
        raw_id = _decision_value(decision, "raw_id")
        status_value = _decision_value(decision, "screening_status")
        reason = _decision_value(decision, "screening_reason")
        if not isinstance(raw_id, str) or not raw_id:
            raise ValueError("screening decision raw_id 不能为空")
        if raw_id in by_id:
            raise ValueError(f"screening decision raw_id 重复: {raw_id}")
        try:
            status = status_value if isinstance(status_value, ScreeningStatus) else ScreeningStatus(status_value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"screening decision 状态无效: {raw_id}") from exc
        if status is ScreeningStatus.DUPLICATE:
            raise ValueError("相关性清洗不得生成 DUPLICATE")
        if not isinstance(reason, str) or not reason.strip():
            raise ValueError(f"screening decision 原因不能为空: {raw_id}")
        by_id[raw_id] = (status, reason.strip())

    expected_ids = [comment.raw_id for comment in imported.dataset.comments]
    if any(raw_id is None for raw_id in expected_ids):
        raise ValueError("原始比赛数据缺少 raw_id")
    expected = {str(raw_id) for raw_id in expected_ids}
    actual = set(by_id)
    if actual != expected:
        missing = sorted(expected - actual)
        extra = sorted(actual - expected)
        raise ValueError(
            "screening decisions 必须完整覆盖原始数据"
            f"；missing={missing[:5]}；extra={extra[:5]}"
        )

    comments = []
    for comment in imported.dataset.comments:
        assert comment.raw_id is not None
        status, reason = by_id[comment.raw_id]
        comments.append(
            comment.model_copy(
                update={"screening_status": status, "screening_reason": reason}
            )
        )
    return CompetitionWorkbookImport(
        dataset=ValidatedDataset(
            comments=comments,
            human_annotations=list(imported.dataset.human_annotations),
            field_mapping=imported.dataset.field_mapping,
            data_health=_data_health(comments),
        ),
        source_manifest=imported.source_manifest,
    )


def load_competition_workbook(workbook_path: str | Path) -> CompetitionWorkbookImport:
    path = Path(workbook_path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("比赛原始语料必须是 .xlsx 工作簿")
    if not path.is_file():
        raise FileNotFoundError(path)

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if WORKSHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"缺少工作表: {WORKSHEET_NAME}")
        sheet = workbook[WORKSHEET_NAME]
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(next(rows, ()))
        if headers != EXPECTED_HEADERS:
            raise ValueError("工作表表头与比赛原始语料合同不一致")

        comments: list[CommentRecord] = []
        for row_number, row in enumerate(rows, start=2):
            if len(row) != len(EXPECTED_HEADERS):
                raise ValueError(f"第 {row_number} 行列数不等于 19")
            values = dict(zip(EXPECTED_HEADERS, row, strict=True))
            raw_id = scalar(values["语料编号"])
            raw_type = scalar(values["样本类型"])
            raw_content = scalar(values["原始内容"])
            source_platform = scalar(values["来源平台"])
            if raw_id is None or raw_type is None or raw_content is None or source_platform is None:
                raise ValueError(f"第 {row_number} 行缺少语料编号、样本类型、原始内容或来源平台")
            if raw_type not in RAW_TYPE_MAPPING:
                raise ValueError(f"第 {row_number} 行出现未知业务类型: {raw_type}")
            if source_platform not in EXPECTED_PLATFORMS:
                raise ValueError(f"第 {row_number} 行出现未知来源平台: {source_platform}")
            original_url, platform_url_available = normalize_original_url(values["来源链接"])
            comments.append(
                CommentRecord(
                    comment_id=raw_id,
                    raw_id=raw_id,
                    sample_type=RAW_TYPE_MAPPING[raw_type],
                    raw_sample_type=raw_type,
                    raw_content=raw_content,
                    context_content=scalar(values["上下文内容"]),
                    source_note=scalar(values["备注"]),
                    source_platform=source_platform,
                    original_url=original_url,
                    platform_url_available=platform_url_available,
                    collected_at=_formatted_date(values["采集日期"]),
                    screening_status=ScreeningStatus.REVIEW,
                    screening_reason="pending relevance screening",
                    source=SourceReference(
                        source_id=raw_id,
                        source_type=SourceType.USER_COMMENT,
                        source_ref=normalized_source_ref(workbook_name=path.name, row_number=row_number),
                    ),
                )
            )
    finally:
        workbook.close()

    if len(comments) != 484:
        raise ValueError(f"工作表记录数必须为 484，实际为 {len(comments)}")
    raw_ids = [comment.raw_id for comment in comments]
    if len(set(raw_ids)) != len(raw_ids):
        raise ValueError("语料编号必须唯一")
    raw_types = {comment.raw_sample_type for comment in comments}
    if raw_types != set(RAW_TYPE_MAPPING):
        raise ValueError("工作表业务类型必须恰好覆盖 10 个已知类型")
    platforms = {comment.source_platform for comment in comments}
    if platforms != EXPECTED_PLATFORMS:
        raise ValueError("工作表来源平台必须恰好覆盖 5 个已知平台")

    dataset = ValidatedDataset(
        comments=comments,
        human_annotations=[],
        field_mapping=_field_mapping(),
        data_health=_data_health(comments),
    )
    return CompetitionWorkbookImport(
        dataset=dataset,
        source_manifest=CompetitionSourceManifest(
            dataset_version=DATASET_VERSION,
            source_sha256=_source_sha256(path),
            imported_at=datetime.now(timezone.utc).isoformat(),
            source_filename=path.name,
            source_sheet=WORKSHEET_NAME,
            record_count=len(comments),
        ),
    )
